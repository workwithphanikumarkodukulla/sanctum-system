"""Native PPTX parser using python-pptx, capturing slides, titles, text, tables, and image metadata."""
from __future__ import annotations

import io
import logging
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, BinaryIO, Sequence
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.evidence.builder import EvidenceBuilder
from app.evidence.schema import EvidenceElement
from app.parsers.base import BaseParser

logger = logging.getLogger(__name__)


@dataclass
class PptxTextItem:
    """Represents a text frame or title on a slide."""
    slide_number: int
    reading_order: int
    text: str
    is_title: bool
    shape_name: str
    shape_id: int | None
    bbox: list[float]  # [left, top, right, bottom]
    is_bullet_list: bool = False


@dataclass
class PptxTableItem:
    """Represents a table structure on a slide."""
    slide_number: int
    reading_order: int
    headers: list[str]
    rows: list[list[str]]
    num_rows: int
    num_cols: int
    shape_name: str
    shape_id: int | None
    bbox: list[float]

    @property
    def text(self) -> str:
        """Markdown formatted representation of the table."""
        if not self.headers and not self.rows:
            return ""
        header_line = " | ".join(self.headers)
        sep_line = " | ".join(["---"] * len(self.headers))
        body_lines = [" | ".join(r) for r in self.rows]
        return f"{header_line}\n{sep_line}\n" + "\n".join(body_lines)


@dataclass
class PptxImageItem:
    """Represents an embedded image on a slide without performing OCR."""
    slide_number: int
    reading_order: int
    shape_name: str
    shape_id: int | None
    content_type: str
    size_bytes: int
    bbox: list[float]
    image_bytes: bytes | None = None


@dataclass
class PptxSlide:
    """Structured representation of a single slide."""
    slide_number: int
    title: str | None
    items: list[PptxTextItem | PptxTableItem | PptxImageItem] = field(default_factory=list)
    text_items: list[PptxTextItem] = field(default_factory=list)
    table_items: list[PptxTableItem] = field(default_factory=list)
    image_items: list[PptxImageItem] = field(default_factory=list)


@dataclass
class PptxParseResult:
    """Structured result of PPTX parsing slide-by-slide."""
    parser_name: str = "pptx_parser"
    source_document: str = ""
    slides: list[PptxSlide] = field(default_factory=list)
    metadata: dict[str, Any] = field(default_factory=dict)
    total_pages: int = 1
    parser_used: str = "python-pptx"

    _cached_elements: list[EvidenceElement] | None = field(default=None, repr=False)
    _doc_id: str = field(default="doc", repr=False)
    _hash: str = field(default="", repr=False)

    def to_evidence_elements(
        self,
        document_id: str | None = None,
        file_hash: str | None = None,
    ) -> list[EvidenceElement]:
        """Normalize structured PPTX elements into canonical EvidenceElements."""
        doc_id = document_id or self._doc_id
        f_hash = file_hash or self._hash
        elements: list[EvidenceElement] = []
        elem_idx = 1

        for slide in self.slides:
            for item in slide.items:
                prov = {
                    "source": "pptx",
                    "document_id": doc_id,
                    "slide": item.slide_number,
                    "shape_name": item.shape_name,
                    "shape_id": item.shape_id,
                    "reading_order": item.reading_order,
                    "bbox": item.bbox,
                    "extraction_model": "python-pptx",
                    "file_hash": f_hash,
                }

                if isinstance(item, PptxTextItem):
                    elem_type = "heading" if item.is_title else ("list" if item.is_bullet_list else "text")
                    elem = EvidenceBuilder.build_element(
                        element_idx=elem_idx,
                        document_id=doc_id,
                        page=item.slide_number,
                        elem_type=elem_type,
                        text=item.text,
                        bbox=item.bbox,
                        confidence=1.0,
                        extraction_model="python-pptx",
                        file_hash=f_hash,
                        slide=item.slide_number,
                        reading_order=item.reading_order,
                        provenance=prov,
                        metadata={
                            "is_title": item.is_title,
                            "is_heading": item.is_title,
                            "is_list": item.is_bullet_list,
                            "is_native_text": True,
                            "is_digital": True,
                            "shape_name": item.shape_name,
                            "shape_id": item.shape_id,
                            "reading_order": item.reading_order,
                            "source": "pptx",
                        },
                    )
                    elements.append(elem)
                    elem_idx += 1

                elif isinstance(item, PptxTableItem):
                    elem = EvidenceBuilder.build_element(
                        element_idx=elem_idx,
                        document_id=doc_id,
                        page=item.slide_number,
                        elem_type="table",
                        text=item.text,
                        table_data={"headers": item.headers, "rows": item.rows},
                        bbox=item.bbox,
                        confidence=1.0,
                        extraction_model="python-pptx",
                        file_hash=f_hash,
                        slide=item.slide_number,
                        reading_order=item.reading_order,
                        provenance=prov,
                        metadata={
                            "shape_name": item.shape_name,
                            "shape_id": item.shape_id,
                            "num_rows": item.num_rows,
                            "num_cols": item.num_cols,
                            "total_rows": item.num_rows,
                            "reading_order": item.reading_order,
                            "is_table": True,
                            "is_digital": True,
                            "source": "pptx_table",
                        },
                    )
                    elements.append(elem)
                    elem_idx += 1

                elif isinstance(item, PptxImageItem):
                    elem = EvidenceBuilder.build_element(
                        element_idx=elem_idx,
                        document_id=doc_id,
                        page=item.slide_number,
                        elem_type="image",
                        text=None,
                        bbox=item.bbox,
                        confidence=1.0,
                        extraction_model="python-pptx",
                        file_hash=f_hash,
                        slide=item.slide_number,
                        reading_order=item.reading_order,
                        provenance=prov,
                        metadata={
                            "shape_name": item.shape_name,
                            "shape_id": item.shape_id,
                            "content_type": item.content_type,
                            "size_bytes": item.size_bytes,
                            "reading_order": item.reading_order,
                            "is_image": True,
                            "source": "pptx_image",
                            "image_bytes": item.image_bytes,
                        },
                    )
                    elements.append(elem)
                    elem_idx += 1

        return elements

    @property
    def elements(self) -> list[EvidenceElement]:
        """Convenience property for pipeline compatibility."""
        if self._cached_elements is None:
            self._cached_elements = self.to_evidence_elements()
        return self._cached_elements


class PptxParser(BaseParser):
    """Native parser for PowerPoint files extracting slide structure in reading order."""

    @classmethod
    def _read_bytes(cls, input_data: bytes | BinaryIO | str | Path) -> bytes:
        if hasattr(input_data, "read"):
            return input_data.read()
        if isinstance(input_data, (str, Path)):
            with open(input_data, "rb") as f:
                return f.read()
        if isinstance(input_data, bytes):
            return input_data
        if isinstance(input_data, (bytearray, memoryview)):
            return bytes(input_data)
        raise TypeError(f"Expected bytes, Path, or file stream, got {type(input_data).__name__}")

    async def parse(
        self,
        file_bytes: bytes | BinaryIO | str | Path,
        filename: str = "",
        document_id: str = "doc",
        file_hash: str = "",
        **kwargs: Any,
    ) -> PptxParseResult:
        """Parse PPTX file bytes into a structured PptxParseResult."""
        raw_bytes = self._read_bytes(file_bytes)
        source_doc = filename or "presentation.pptx"

        if not raw_bytes:
            return PptxParseResult(
                parser_name="pptx_parser",
                source_document=source_doc,
                slides=[],
                metadata={"total_slides": 0, "is_empty": True},
                total_pages=1,
                parser_used="python-pptx",
                _doc_id=document_id,
                _hash=file_hash,
            )

        try:
            prs = Presentation(io.BytesIO(raw_bytes))
        except Exception as exc:
            logger.warning("Failed to parse PPTX presentation '%s': %s", source_doc, exc)
            return PptxParseResult(
                parser_name="pptx_parser",
                source_document=source_doc,
                slides=[],
                metadata={"total_slides": 0, "is_corrupt": True, "error": str(exc)},
                total_pages=1,
                parser_used="python-pptx",
                _doc_id=document_id,
                _hash=file_hash,
            )

        total_slides = len(prs.slides)
        slide_objects: list[PptxSlide] = []

        total_texts = 0
        total_tables = 0
        total_images = 0

        for slide_idx, slide in enumerate(prs.slides):
            slide_num = slide_idx + 1
            reading_order = 0

            # 1. Identify slide title
            title_text = None
            if slide.shapes.title and slide.shapes.title.has_text_frame:
                title_text = slide.shapes.title.text.strip()

            slide_obj = PptxSlide(slide_number=slide_num, title=title_text)

            # 2. Sort shapes top-to-bottom, left-to-right for consistent reading order
            sorted_shapes = sorted(
                slide.shapes,
                key=lambda s: (getattr(s, "top", 0) or 0, getattr(s, "left", 0) or 0),
            )

            for shape in sorted_shapes:
                left = float(getattr(shape, "left", 0) or 0)
                top = float(getattr(shape, "top", 0) or 0)
                width = float(getattr(shape, "width", 0) or 0)
                height = float(getattr(shape, "height", 0) or 0)
                bbox = [left, top, left + width, top + height]

                # Check for table
                if shape.has_table:
                    tbl = shape.table
                    rows_data: list[list[str]] = []
                    for row in tbl.rows:
                        row_cells = [cell.text.strip() for cell in row.cells]
                        rows_data.append(row_cells)

                    if rows_data:
                        headers = rows_data[0]
                        body_rows = rows_data[1:] if len(rows_data) > 1 else []
                        reading_order += 1
                        total_tables += 1

                        table_item = PptxTableItem(
                            slide_number=slide_num,
                            reading_order=reading_order,
                            headers=headers,
                            rows=body_rows,
                            num_rows=len(rows_data),
                            num_cols=len(headers),
                            shape_name=shape.name,
                            shape_id=getattr(shape, "shape_id", None),
                            bbox=bbox,
                        )
                        slide_obj.items.append(table_item)
                        slide_obj.table_items.append(table_item)

                # Check for picture / image
                elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE or hasattr(shape, "image"):
                    reading_order += 1
                    total_images += 1
                    content_type = "image/unknown"
                    size_bytes = 0
                    img_blob = None
                    if hasattr(shape, "image"):
                        content_type = getattr(shape.image, "content_type", "image/unknown")
                        img_blob = getattr(shape.image, "blob", None)
                        size_bytes = len(img_blob) if img_blob else 0

                    image_item = PptxImageItem(
                        slide_number=slide_num,
                        reading_order=reading_order,
                        shape_name=shape.name,
                        shape_id=getattr(shape, "shape_id", None),
                        content_type=content_type,
                        size_bytes=size_bytes,
                        bbox=bbox,
                        image_bytes=img_blob,
                    )
                    slide_obj.items.append(image_item)
                    slide_obj.image_items.append(image_item)

                # Check for text frame / text box / title
                elif shape.has_text_frame:
                    text = shape.text.strip()
                    if text:
                        reading_order += 1
                        total_texts += 1
                        is_title = (shape == slide.shapes.title)
                        is_bullet_list = any(
                            (getattr(p, "level", 0) or 0) > 0 or p.text.strip().startswith(("•", "-", "*", "–", "—"))
                            for p in shape.text_frame.paragraphs
                            if p.text.strip()
                        )

                        text_item = PptxTextItem(
                            slide_number=slide_num,
                            reading_order=reading_order,
                            text=text,
                            is_title=is_title,
                            shape_name=shape.name,
                            shape_id=getattr(shape, "shape_id", None),
                            bbox=bbox,
                            is_bullet_list=is_bullet_list,
                        )
                        slide_obj.items.append(text_item)
                        slide_obj.text_items.append(text_item)

            slide_objects.append(slide_obj)

        meta = {
            "source_document": source_doc,
            "total_slides": total_slides,
            "total_texts": total_texts,
            "total_tables": total_tables,
            "total_images": total_images,
            "is_empty": total_slides == 0,
        }

        return PptxParseResult(
            parser_name="pptx_parser",
            source_document=source_doc,
            slides=slide_objects,
            metadata=meta,
            total_pages=max(1, total_slides),
            parser_used="python-pptx",
            _doc_id=document_id,
            _hash=file_hash,
        )


pptx_parser = PptxParser()
