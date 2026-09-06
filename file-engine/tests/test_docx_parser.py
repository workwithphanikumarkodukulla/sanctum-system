"""Unit tests for the DocxParser covering reading order, headings, paragraphs, tables, and empty docs."""
from __future__ import annotations

import io
import pytest
from docx import Document

from app.parsers.docx import DocxParagraph, DocxParseResult, DocxTable, docx_parser


def create_docx_bytes(build_fn) -> bytes:
    doc = Document()
    build_fn(doc)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.mark.anyio
async def test_empty_docx_file():
    """Verify empty input returns empty DocxParseResult cleanly."""
    # 0-byte input
    res_zero = await docx_parser.parse(b"", filename="empty.docx")
    assert isinstance(res_zero, DocxParseResult)
    assert res_zero.parser_name == "docx_parser"
    assert len(res_zero.items) == 0
    assert res_zero.metadata["is_empty"] is True

    # Empty Document without paragraphs or tables
    empty_doc_bytes = create_docx_bytes(lambda doc: None)
    res_empty_doc = await docx_parser.parse(empty_doc_bytes, filename="blank.docx")
    assert isinstance(res_empty_doc, DocxParseResult)
    assert len(res_empty_doc.items) == 0
    assert res_empty_doc.metadata["is_empty"] is True


@pytest.mark.anyio
async def test_docx_only_paragraphs():
    """Verify DOCX containing only regular paragraphs."""
    def build(doc):
        doc.add_paragraph("First paragraph describing inspection procedures.")
        doc.add_paragraph("Second paragraph detailing safety measures.")
        doc.add_paragraph("Third paragraph summarizing findings.")

    docx_bytes = create_docx_bytes(build)
    res = await docx_parser.parse(docx_bytes, filename="procedures.docx")

    assert isinstance(res, DocxParseResult)
    assert len(res.items) == 3
    assert len(res.paragraphs) == 3
    assert len(res.tables) == 0
    assert res.metadata["total_paragraphs"] == 3

    for idx, p in enumerate(res.paragraphs, start=1):
        assert isinstance(p, DocxParagraph)
        assert p.index == idx
        assert p.reading_order == idx
        assert p.is_heading is False
        assert p.source_document == "procedures.docx"

    assert res.paragraphs[0].text == "First paragraph describing inspection procedures."
    assert res.paragraphs[2].text == "Third paragraph summarizing findings."


@pytest.mark.anyio
async def test_docx_headings_and_paragraphs():
    """Verify DOCX with Title, Heading 1, Heading 2, and text paragraphs."""
    def build(doc):
        doc.add_heading("Pressure Vessel Technical Manual", level=0)
        doc.add_paragraph("Preliminary safety notice.")
        doc.add_heading("Section 1: Operating Thresholds", level=1)
        doc.add_paragraph("Maximum allowable working pressure is 1500 kPa.")
        doc.add_heading("Subsection 1.1: Temperature Limits", level=2)
        doc.add_paragraph("Operating range: -20 C to 180 C.")

    docx_bytes = create_docx_bytes(build)
    res = await docx_parser.parse(docx_bytes, filename="manual.docx")

    assert len(res.paragraphs) == 6
    assert res.metadata["total_paragraphs"] == 6

    # Title (level 0)
    p0 = res.paragraphs[0]
    assert p0.is_heading is True
    assert p0.heading_level == 0
    assert "Pressure Vessel Technical Manual" in p0.text

    # Heading 1
    p2 = res.paragraphs[2]
    assert p2.is_heading is True
    assert p2.heading_level == 1
    assert "Section 1: Operating Thresholds" in p2.text

    # Paragraph under Heading 1
    p3 = res.paragraphs[3]
    assert p3.is_heading is False
    assert p3.heading_level is None
    assert "Maximum allowable working pressure" in p3.text

    # Heading 2
    p4 = res.paragraphs[4]
    assert p4.is_heading is True
    assert p4.heading_level == 2
    assert "Subsection 1.1" in p4.text


@pytest.mark.anyio
async def test_docx_with_tables():
    """Verify DOCX table extraction preserving row order, column order, and cell values."""
    def build(doc):
        table = doc.add_table(rows=3, cols=3)
        data = [
            ["Vessel ID", "Design Temp (C)", "Test Pressure (bar)"],
            ["V-101", "250", "32.5"],
            ["V-102", "300", "45.0"],
        ]
        for r_idx, row in enumerate(data):
            for c_idx, val in enumerate(row):
                table.cell(r_idx, c_idx).text = val

    docx_bytes = create_docx_bytes(build)
    res = await docx_parser.parse(docx_bytes, filename="specs.docx")

    assert len(res.tables) == 1
    table_item = res.tables[0]
    assert isinstance(table_item, DocxTable)
    assert table_item.index == 1
    assert table_item.num_rows == 3
    assert table_item.num_cols == 3
    assert table_item.headers == ["Vessel ID", "Design Temp (C)", "Test Pressure (bar)"]
    assert table_item.rows == [
        ["V-101", "250", "32.5"],
        ["V-102", "300", "45.0"],
    ]
    # Check cell values
    assert table_item.rows[0][0] == "V-101"
    assert table_item.rows[0][2] == "32.5"
    assert table_item.rows[1][0] == "V-102"
    assert table_item.rows[1][2] == "45.0"


@pytest.mark.anyio
async def test_docx_interleaved_reading_order():
    """Verify exact interleaved reading order: Heading -> Para -> Para -> Table -> Para."""
    def build(doc):
        doc.add_heading("Equipment Summary", level=1)  # order 1
        doc.add_paragraph("Introductory text line 1.")  # order 2
        doc.add_paragraph("Introductory text line 2.")  # order 3
        tbl = doc.add_table(rows=2, cols=2)             # order 4
        tbl.cell(0, 0).text = "Col A"
        tbl.cell(0, 1).text = "Col B"
        tbl.cell(1, 0).text = "Val 1"
        tbl.cell(1, 1).text = "Val 2"
        doc.add_paragraph("Post-table trailing comments.")  # order 5

    docx_bytes = create_docx_bytes(build)
    res = await docx_parser.parse(docx_bytes, filename="report.docx")

    assert len(res.items) == 5
    assert isinstance(res.items[0], DocxParagraph)
    assert res.items[0].reading_order == 1
    assert res.items[0].is_heading is True

    assert isinstance(res.items[1], DocxParagraph)
    assert res.items[1].reading_order == 2

    assert isinstance(res.items[2], DocxParagraph)
    assert res.items[2].reading_order == 3

    assert isinstance(res.items[3], DocxTable)
    assert res.items[3].reading_order == 4
    assert res.items[3].headers == ["Col A", "Col B"]

    assert isinstance(res.items[4], DocxParagraph)
    assert res.items[4].reading_order == 5
    assert "Post-table trailing comments" in res.items[4].text

    # Normalization helper verification
    elements = res.to_evidence_elements(document_id="doc_test_norm", file_hash="hash999")
    assert len(elements) == 5
    assert elements[0].type == "text"
    assert elements[3].type == "table"
    assert elements[3].table_data["headers"] == ["Col A", "Col B"]
