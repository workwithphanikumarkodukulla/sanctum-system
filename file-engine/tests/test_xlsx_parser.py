"""Unit tests for the XlsxParser covering normal worksheets, multiple sheets, formulas, merged cells, and empty cells."""
from __future__ import annotations

import io
import openpyxl
import pytest

from app.parsers.xlsx import (
    XlsxCell,
    XlsxParseResult,
    XlsxWorksheet,
    xlsx_parser,
)


def create_xlsx_bytes(build_fn) -> bytes:
    wb = openpyxl.Workbook()
    build_fn(wb)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.mark.anyio
async def test_xlsx_normal_worksheet():
    """Verify single worksheet with headers and rows preserving order."""
    def build(wb):
        ws = wb.active
        ws.title = "Inventory"
        ws.append(["Part_ID", "Description", "Stock_Qty", "Unit_Price"])
        ws.append(["P-101", "Ball Bearing", 50, 12.50])
        ws.append(["P-102", "O-Ring Seal", 120, 2.75])

    raw_bytes = create_xlsx_bytes(build)
    res: XlsxParseResult = await xlsx_parser.parse(raw_bytes, filename="inventory.xlsx")

    assert isinstance(res, XlsxParseResult)
    assert res.parser_name == "xlsx_parser"
    assert len(res.worksheets) == 1
    ws = res.worksheets[0]
    assert ws.name == "Inventory"
    assert ws.headers == ["Part_ID", "Description", "Stock_Qty", "Unit_Price"]
    assert len(ws.data_rows) == 2
    assert ws.data_rows[0] == ["P-101", "Ball Bearing", "50", "12.5"]
    assert ws.data_rows[1] == ["P-102", "O-Ring Seal", "120", "2.75"]


@pytest.mark.anyio
async def test_xlsx_multiple_worksheets():
    """Verify multiple worksheets are captured with their sheet names and indices."""
    def build(wb):
        ws1 = wb.active
        ws1.title = "Q1_Summary"
        ws1.append(["Metric", "Value"])
        ws1.append(["Revenue", "100000"])

        ws2 = wb.create_sheet(title="Q2_Summary")
        ws2.append(["Metric", "Value"])
        ws2.append(["Revenue", "120000"])

        ws3 = wb.create_sheet(title="Settings")
        ws3.append(["Key", "Setting"])
        ws3.append(["Currency", "USD"])

    raw_bytes = create_xlsx_bytes(build)
    res = await xlsx_parser.parse(raw_bytes, filename="finances.xlsx")

    assert len(res.worksheets) == 3
    assert res.metadata["total_sheets"] == 3
    sheet_names = [w.name for w in res.worksheets]
    assert sheet_names == ["Q1_Summary", "Q2_Summary", "Settings"]
    assert res.worksheets[0].sheet_index == 1
    assert res.worksheets[1].sheet_index == 2
    assert res.worksheets[2].sheet_index == 3


@pytest.mark.anyio
async def test_xlsx_formulas_preserved_verbatim():
    """Verify formulas are NOT evaluated and original formula string is preserved verbatim."""
    def build(wb):
        ws = wb.active
        ws.title = "Calculations"
        ws.append(["Value A", "Value B", "Sum Formula", "Product Formula"])
        ws.append([10, 20, "=A2+B2", "=A2*B2"])
        ws.append([30, 40, "=SUM(A2:B3)", "=AVERAGE(A2:B3)"])

    raw_bytes = create_xlsx_bytes(build)
    res = await xlsx_parser.parse(raw_bytes, filename="calc.xlsx")

    ws = res.worksheets[0]
    assert len(ws.formula_cells) == 4

    formula_strings = {fc.formula for fc in ws.formula_cells}
    assert "=A2+B2" in formula_strings
    assert "=A2*B2" in formula_strings
    assert "=SUM(A2:B3)" in formula_strings
    assert "=AVERAGE(A2:B3)" in formula_strings

    # Verify specific cell coordinates
    c_sum = next(fc for fc in ws.formula_cells if fc.coordinate == "C2")
    assert c_sum.is_formula is True
    assert c_sum.formula == "=A2+B2"
    assert c_sum.row == 2
    assert c_sum.column == 3

    # Check normalization to EvidenceElement
    elements = res.to_evidence_elements(document_id="doc_formula_test", file_hash="h1")
    formula_elems = [e for e in elements if e.type == "formula"]
    assert len(formula_elems) == 4
    assert formula_elems[0].text == "=A2+B2"
    assert formula_elems[0].row == 2
    assert formula_elems[0].column == 3
    assert formula_elems[0].metadata["cell_coordinate"] == "C2"


@pytest.mark.anyio
async def test_xlsx_merged_cells():
    """Verify merged cell ranges are detected and preserved in metadata."""
    def build(wb):
        ws = wb.active
        ws.title = "HeaderTable"
        ws.append(["Quarterly Department Report", "", "", ""])
        ws.merge_cells("A1:D1")
        ws.append(["Dept", "Lead", "Budget", "Spent"])
        ws.append(["Engineering", "Alice", 50000, 42000])
        ws.merge_cells("A4:B4")
        ws["A4"] = "Totals / Summary"

    raw_bytes = create_xlsx_bytes(build)
    res = await xlsx_parser.parse(raw_bytes, filename="report.xlsx")

    ws = res.worksheets[0]
    assert len(ws.merged_cells) == 2
    assert any("A1:D1" in r for r in ws.merged_cells)
    assert any("A4:B4" in r for r in ws.merged_cells)
    assert res.metadata["total_merged_ranges"] == 2


@pytest.mark.anyio
async def test_xlsx_empty_cells_and_empty_sheet():
    """Verify empty cells within rows and completely empty sheets are distinguished cleanly."""
    def build(wb):
        ws1 = wb.active
        ws1.title = "SparseData"
        # Row 1: has holes
        ws1["A1"] = "Item"
        ws1["C1"] = "Category"  # B1 is empty
        # Row 2:
        ws1["A2"] = "Motor"
        ws1["B2"] = None         # explicitly empty
        ws1["C2"] = "Electrical"

        # Sheet 2: completely blank
        wb.create_sheet(title="BlankSheet")

    raw_bytes = create_xlsx_bytes(build)
    res = await xlsx_parser.parse(raw_bytes, filename="sparse.xlsx")

    # Check SparseData sheet
    ws_sparse = res.worksheets[0]
    row1 = ws_sparse.rows[0]
    cell_b1 = next(c for c in row1.cells if c.coordinate == "B1")
    assert cell_b1.is_empty is True
    assert cell_b1.value_str == ""

    # Check BlankSheet
    ws_blank = res.worksheets[1]
    assert len(ws_blank.rows) == 0
    assert ws_blank.headers == []
    assert ws_blank.data_rows == []
