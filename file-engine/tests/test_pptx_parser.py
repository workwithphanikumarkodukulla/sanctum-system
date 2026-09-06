"""Unit tests for the PptxParser covering titles, text, multiple slides, tables, and images."""
from __future__ import annotations

import io
import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from app.parsers.pptx import (
    PptxImageItem,
    PptxParseResult,
    PptxTableItem,
    PptxTextItem,
    pptx_parser,
)


def create_sample_png_bytes() -> bytes:
    """Helper to generate a tiny valid PNG image."""
    img = Image.new("RGB", (60, 30), color=(100, 150, 200))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


@pytest.mark.anyio
async def test_pptx_empty_input():
    """Verify empty input handling."""
    res = await pptx_parser.parse(b"", filename="empty.pptx")
    assert isinstance(res, PptxParseResult)
    assert res.parser_name == "pptx_parser"
    assert len(res.slides) == 0
    assert res.metadata["is_empty"] is True


@pytest.mark.anyio
async def test_pptx_title_and_text():
    """Verify single slide with Title and Body text box."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Q4 Operational Review"
    slide.placeholders[1].text = "Sanctum Engineering Operations"

    buf = io.BytesIO()
    prs.save(buf)

    res = await pptx_parser.parse(buf.getvalue(), filename="review.pptx")

    assert isinstance(res, PptxParseResult)
    assert len(res.slides) == 1
    s1 = res.slides[0]
    assert s1.slide_number == 1
    assert s1.title == "Q4 Operational Review"

    text_contents = [t.text for t in s1.text_items]
    assert "Q4 Operational Review" in text_contents
    assert "Sanctum Engineering Operations" in text_contents

    title_item = next(t for t in s1.text_items if t.is_title)
    assert title_item.text == "Q4 Operational Review"
    assert len(title_item.bbox) == 4


@pytest.mark.anyio
async def test_pptx_multiple_slides():
    """Verify presentation with multiple slides preserving slide numbers and reading order."""
    prs = Presentation()

    # Slide 1
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "Agenda & Objectives"
    tb1 = s1.shapes.add_textbox(Inches(1), Inches(2), Inches(5), Inches(1))
    tb1.text_frame.text = "1. Facility Audit\n2. Pressure Check"

    # Slide 2
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "Audit Findings"
    tb2 = s2.shapes.add_textbox(Inches(1), Inches(2), Inches(5), Inches(1))
    tb2.text_frame.text = "All primary valves tested successfully."

    buf = io.BytesIO()
    prs.save(buf)

    res = await pptx_parser.parse(buf.getvalue(), filename="agenda.pptx")

    assert len(res.slides) == 2
    assert res.metadata["total_slides"] == 2
    assert res.slides[0].slide_number == 1
    assert res.slides[0].title == "Agenda & Objectives"
    assert res.slides[1].slide_number == 2
    assert res.slides[1].title == "Audit Findings"


@pytest.mark.anyio
async def test_pptx_with_table():
    """Verify slide with table captures headers, rows, columns, and shape metadata."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    table_shape = slide.shapes.add_table(3, 2, Inches(1), Inches(1), Inches(5), Inches(3))
    tbl = table_shape.table
    tbl.cell(0, 0).text = "Sensor ID"
    tbl.cell(0, 1).text = "Reading (bar)"
    tbl.cell(1, 0).text = "SN-01"
    tbl.cell(1, 1).text = "14.2"
    tbl.cell(2, 0).text = "SN-02"
    tbl.cell(2, 1).text = "16.8"

    buf = io.BytesIO()
    prs.save(buf)

    res = await pptx_parser.parse(buf.getvalue(), filename="sensors.pptx")

    assert len(res.slides) == 1
    assert len(res.slides[0].table_items) == 1
    table_item: PptxTableItem = res.slides[0].table_items[0]

    assert table_item.slide_number == 1
    assert table_item.num_rows == 3
    assert table_item.num_cols == 2
    assert table_item.headers == ["Sensor ID", "Reading (bar)"]
    assert table_item.rows == [
        ["SN-01", "14.2"],
        ["SN-02", "16.8"],
    ]
    assert "Sensor ID | Reading (bar)" in table_item.text


@pytest.mark.anyio
async def test_pptx_with_image_shape():
    """Verify slide with embedded picture identifies image metadata without running OCR."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    png_bytes = create_sample_png_bytes()
    png_stream = io.BytesIO(png_bytes)

    slide.shapes.add_picture(png_stream, Inches(1), Inches(1), width=Inches(2), height=Inches(1))

    buf = io.BytesIO()
    prs.save(buf)

    res = await pptx_parser.parse(buf.getvalue(), filename="diagram.pptx")

    assert len(res.slides) == 1
    assert len(res.slides[0].image_items) == 1
    img_item: PptxImageItem = res.slides[0].image_items[0]

    assert img_item.slide_number == 1
    assert "image" in img_item.content_type.lower()
    assert img_item.size_bytes > 0
    assert len(img_item.bbox) == 4

    # Verify normalization produces an image element without OCR error
    elements = res.to_evidence_elements(document_id="doc_img_test", file_hash="hash_img")
    assert len(elements) == 1
    assert elements[0].type == "image"
    assert elements[0].slide == 1
    assert elements[0].confidence == 1.0
