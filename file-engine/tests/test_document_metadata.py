"""Unit tests for the canonical DocumentMetadata and WorksheetInfo schema models and API integration."""
from __future__ import annotations

import io
import pytest
from starlette.testclient import TestClient

from app.evidence.schema import DocumentEvidence, DocumentMetadata, EvidenceElement, WorksheetInfo
from app.evidence.storage import storage
from app.main import app
from app.processing.pipeline import pipeline

client = TestClient(app)


def test_worksheet_info_model():
    """Verify WorksheetInfo accurately captures spreadsheet sheet characteristics."""
    ws = WorksheetInfo(
        name="Q1_Revenue",
        index=1,
        row_count=100,
        column_count=12,
        has_formulas=True,
        merged_ranges=["A1:D1", "F2:G2"],
    )

    assert ws.name == "Q1_Revenue"
    assert ws.index == 1
    assert ws.row_count == 100
    assert ws.column_count == 12
    assert ws.has_formulas is True
    assert ws.merged_ranges == ["A1:D1", "F2:G2"]


def test_document_metadata_model_instantiation():
    """Verify DocumentMetadata captures document_id, filename, hashes, counts, and processing stats."""
    ws = WorksheetInfo(name="Sheet1", index=1, row_count=50, column_count=5)
    meta = DocumentMetadata(
        document_id="doc_meta_001",
        filename="financial_deck.pptx",
        file_type="pptx",
        mime_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        file_size=1048576,
        file_hash="e3b0c44298fc1c149afbf4c8996fb92427ae41e4649b934ca495991b7852b855",
        processing_status="completed",
        parser_used="python-pptx",
        total_pages=12,
        total_slides=12,
        worksheets=[ws],
        created_at="2026-09-05T00:00:00Z",
        completed_at="2026-09-05T00:00:02Z",
        processing_time_ms=1840.5,
        metadata={"author": "Sanctum", "department": "Finance"},
    )

    assert meta.document_id == "doc_meta_001"
    assert meta.filename == "financial_deck.pptx"
    assert meta.file_type == "pptx"
    assert meta.file_size == 1048576
    assert meta.total_slides == 12
    assert len(meta.worksheets) == 1
    assert meta.worksheets[0].name == "Sheet1"
    assert meta.processing_time_ms == 1840.5
    assert meta.metadata["author"] == "Sanctum"


def test_document_evidence_inherits_document_metadata_cleanly():
    """Verify DocumentEvidence is a subclass of DocumentMetadata without duplicating fields."""
    assert issubclass(DocumentEvidence, DocumentMetadata)

    elem = EvidenceElement(
        id="p1_e1",
        document_id="doc_ev_001",
        page=1,
        type="text",
        text="Sample extracted evidence line",
        confidence=1.0,
        extraction_model="test",
        file_hash="hash_123",
        timestamp="2026-09-05T00:00:00Z",
    )

    doc_ev = DocumentEvidence(
        document_id="doc_ev_001",
        filename="report.docx",
        file_type="docx",
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        file_size=2048,
        file_hash="hash_123",
        parser_used="python-docx",
        total_pages=1,
        elements=[elem],
        overall_confidence=1.0,
        processing_status="completed",
        created_at="2026-09-05T00:00:00Z",
        completed_at="2026-09-05T00:00:01Z",
        processing_time_ms=850.25,
    )

    # Verify inheritance
    assert isinstance(doc_ev, DocumentMetadata)
    assert doc_ev.filename == "report.docx"
    assert doc_ev.processing_time_ms == 850.25

    # Extract clean DocumentMetadata without elements
    meta = doc_ev.to_metadata()
    assert isinstance(meta, DocumentMetadata)
    assert not hasattr(meta, "elements")
    assert not hasattr(meta, "overall_confidence")
    assert meta.document_id == "doc_ev_001"
    assert meta.filename == "report.docx"

    # Roundtrip from metadata
    reconstructed = DocumentEvidence.from_metadata(meta, elements=[elem], overall_confidence=1.0)
    assert reconstructed.document_id == doc_ev.document_id
    assert reconstructed.elements[0].text == "Sample extracted evidence line"


def test_backward_compatibility_with_minimal_document_evidence():
    """Verify that existing legacy calls without new metadata fields continue to instantiate properly."""
    doc = DocumentEvidence(
        document_id="doc_legacy",
        total_pages=3,
        elements=[],
        overall_confidence=0.95,
        processing_status="completed",
    )

    assert doc.document_id == "doc_legacy"
    assert doc.total_pages == 3
    assert doc.total_slides is None
    assert doc.worksheets is None
    assert doc.created_at is None
    assert doc.completed_at is None
    assert doc.processing_status == "completed"


@pytest.mark.anyio
async def test_pipeline_populates_metadata_for_pptx():
    """Verify pipeline execution on PPTX populates total_slides, timestamps, and elapsed time."""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Metadata Presentation Test"
    buf = io.BytesIO()
    prs.save(buf)
    pptx_bytes = buf.getvalue()

    doc_id = "doc_meta_pptx_test"
    result = await pipeline.process_document(
        document_id=doc_id,
        file_bytes=pptx_bytes,
        filename="presentation_meta.pptx",
    )

    assert result.document_id == doc_id
    assert result.file_type == "pptx"
    assert result.total_slides == 1
    assert result.created_at is not None
    assert result.completed_at is not None
    assert result.processing_time_ms is not None
    assert result.processing_time_ms >= 0.0


@pytest.mark.anyio
async def test_pipeline_populates_metadata_for_xlsx():
    """Verify pipeline execution on XLSX populates worksheet information models."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SalesData"
    ws.append(["Category", "Revenue"])
    ws.append(["Software", 50000])
    ws["C3"] = "=SUM(B2:B2)"
    buf = io.BytesIO()
    wb.save(buf)
    xlsx_bytes = buf.getvalue()

    doc_id = "doc_meta_xlsx_test"
    result = await pipeline.process_document(
        document_id=doc_id,
        file_bytes=xlsx_bytes,
        filename="sales_meta.xlsx",
    )

    assert result.document_id == doc_id
    assert result.file_type == "xlsx"
    assert result.worksheets is not None
    assert len(result.worksheets) == 1
    ws_info = result.worksheets[0]
    assert ws_info.name == "SalesData"
    assert ws_info.row_count >= 2
    assert ws_info.has_formulas is True


def test_api_document_metadata_endpoint():
    """Verify GET /api/documents/{document_id}/metadata returns canonical DocumentMetadata."""
    doc = DocumentEvidence(
        document_id="doc_api_meta_001",
        filename="annual_report.pdf",
        file_type="pdf",
        mime_type="application/pdf",
        file_size=54321,
        file_hash="hash_pdf_meta",
        parser_used="PyMuPDF",
        total_pages=5,
        elements=[],
        overall_confidence=0.98,
        processing_status="completed",
        created_at="2026-09-05T01:00:00Z",
        completed_at="2026-09-05T01:00:03Z",
        processing_time_ms=2980.0,
    )
    storage.save(doc)

    res = client.get("/api/documents/doc_api_meta_001/metadata")
    assert res.status_code == 200
    data = res.json()

    # Validate against DocumentMetadata schema
    meta = DocumentMetadata(**data)
    assert meta.document_id == "doc_api_meta_001"
    assert meta.filename == "annual_report.pdf"
    assert meta.file_type == "pdf"
    assert meta.total_pages == 5
    assert meta.processing_time_ms == 2980.0
    # Must NOT return elements array in metadata endpoint
    assert "elements" not in data
