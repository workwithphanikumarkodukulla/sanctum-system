"""Local document generation helpers used by the agent tools."""
from __future__ import annotations

from pathlib import Path
from typing import Any


class DocumentGenerator:
    """Generate documents beneath the currently selected workspace root."""

    def __init__(self, root_dir: str | Path) -> None:
        self.set_root_dir(root_dir)

    def set_root_dir(self, root_dir: str | Path) -> None:
        self.root_dir = Path(root_dir).expanduser().resolve()
        self.root_dir.mkdir(parents=True, exist_ok=True)

    def _resolve_path(self, filepath: str, allowed_extensions: tuple[str, ...] | None = None) -> Path:
        if filepath is None or not str(filepath).strip():
            raise ValueError("Output path cannot be empty or whitespace.")
        if "\x00" in str(filepath):
            raise ValueError("Output path contains invalid null byte.")
        try:
            target = (self.root_dir / filepath).expanduser().resolve()
            target.relative_to(self.root_dir)
        except ValueError as exc:
            raise ValueError("Output path escapes the configured workspace root.") from exc
        if allowed_extensions and target.suffix.lower() not in allowed_extensions:
            raise ValueError(
                f"Invalid file extension '{target.suffix}'. Allowed extensions: {', '.join(allowed_extensions)}"
            )
        target.parent.mkdir(parents=True, exist_ok=True)
        return target

    def _relative(self, path: Path) -> str:
        return path.relative_to(self.root_dir).as_posix()

    def generate_excel(self, filepath: str, sheets: list[dict[str, Any]]) -> str:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill

        output = self._resolve_path(filepath, allowed_extensions=(".xlsx", ".xlsm"))

        workbook = Workbook()
        workbook.remove(workbook.active)
        header_fill = PatternFill("solid", fgColor="163B52")
        total_fill = PatternFill("solid", fgColor="D8EEF0")
        for index, sheet_data in enumerate(sheets or [{"title": "Sheet 1", "headers": [], "rows": []}]):
            title = str(sheet_data.get("title") or f"Sheet {index + 1}")[:31]
            sheet = workbook.create_sheet(title)
            headers = list(sheet_data.get("headers") or [])
            if headers:
                sheet.append(headers)
                for cell in sheet[1]:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = header_fill
                    cell.alignment = Alignment(horizontal="center")
            for row in sheet_data.get("rows") or []:
                sheet.append(list(row))
            totals = sheet_data.get("totals_row")
            if totals:
                sheet.append(list(totals))
                for cell in sheet[sheet.max_row]:
                    cell.font = Font(bold=True)
                    cell.fill = total_fill
            sheet.freeze_panes = "A2" if headers else "A1"
            for column in sheet.columns:
                values = [len(str(cell.value or "")) for cell in column]
                if values:
                    sheet.column_dimensions[column[0].column_letter].width = min(max(max(values) + 2, 10), 42)
        workbook.save(output)
        return self._relative(output)

    def generate_presentation(self, filepath: str, title: str, subtitle: str, slides_data: list[dict[str, Any]]) -> str:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.util import Inches, Pt

        output = self._resolve_path(filepath, allowed_extensions=(".pptx",))
        presentation = Presentation()
        presentation.slide_width = Inches(13.333)
        presentation.slide_height = Inches(7.5)
        blank = presentation.slide_layouts[6]

        cover = presentation.slides.add_slide(blank)
        background = cover.background.fill
        background.solid()
        background.fore_color.rgb = RGBColor(13, 25, 35)
        title_box = cover.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(11.7), Inches(1.2))
        title_frame = title_box.text_frame
        title_frame.text = title
        title_frame.paragraphs[0].font.size = Pt(34)
        title_frame.paragraphs[0].font.bold = True
        title_frame.paragraphs[0].font.color.rgb = RGBColor(230, 247, 248)
        subtitle_box = cover.shapes.add_textbox(Inches(0.85), Inches(3.35), Inches(10.8), Inches(0.7))
        subtitle_box.text_frame.text = subtitle or ""
        subtitle_box.text_frame.paragraphs[0].font.size = Pt(16)
        subtitle_box.text_frame.paragraphs[0].font.color.rgb = RGBColor(120, 196, 196)

        for slide_data in slides_data or []:
            slide = presentation.slides.add_slide(blank)
            heading = slide.shapes.add_textbox(Inches(0.7), Inches(0.45), Inches(12), Inches(0.7))
            slide_title = slide_data.get("title") or slide_data.get("heading") or "Untitled"
            heading.text_frame.text = str(slide_title)
            heading.text_frame.paragraphs[0].font.size = Pt(25)
            heading.text_frame.paragraphs[0].font.bold = True
            heading.text_frame.paragraphs[0].font.color.rgb = RGBColor(24, 55, 72)
            cards = slide_data.get("cards") or []
            if cards:
                width = 5.7 if len(cards) <= 2 else 3.8
                for index, card in enumerate(cards):
                    x = 0.7 + index * (width + 0.25)
                    shape = slide.shapes.add_shape(1, Inches(x), Inches(1.6), Inches(width), Inches(3.2))
                    shape.fill.solid()
                    shape.fill.fore_color.rgb = RGBColor(232, 244, 246)
                    shape.line.color.rgb = RGBColor(112, 180, 184)
                    frame = shape.text_frame
                    frame.text = str(card.get("title") or "")
                    frame.paragraphs[0].font.size = Pt(17)
                    frame.paragraphs[0].font.bold = True
                    frame.add_paragraph().text = str(card.get("text") or "")
                    frame.paragraphs[1].font.size = Pt(13)
            else:
                body = slide.shapes.add_textbox(Inches(1.0), Inches(1.7), Inches(11), Inches(4.8))
                frame = body.text_frame
                bullets = slide_data.get("bullet_points")
                if bullets is None:
                    content_val = slide_data.get("content")
                    if isinstance(content_val, list):
                        bullets = []
                        for item in content_val:
                            if isinstance(item, dict) and "headers" in item and "rows" in item:
                                bullets.append(" | ".join(str(h) for h in item["headers"]))
                                for r in item.get("rows", []):
                                    bullets.append(" | ".join(str(c) for c in r))
                            else:
                                bullets.append(str(item))
                    elif content_val:
                        bullets = [str(content_val)]
                    else:
                        bullets = [""]
                for index, point in enumerate(bullets):
                    paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
                    paragraph.text = str(point)
                    paragraph.font.size = Pt(18 if len(bullets) > 4 else 20)
                    paragraph.level = 0
        presentation.save(output)
        return self._relative(output)

    def generate_docx(self, filepath: str, title: str, subtitle: str, sections: list[dict[str, Any]]) -> str:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches, Pt

        output = self._resolve_path(filepath, allowed_extensions=(".docx",))
        document = Document()
        section = document.sections[0]
        section.top_margin = Inches(0.65)
        section.bottom_margin = Inches(0.65)
        heading = document.add_heading(title, 0)
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle_paragraph = document.add_paragraph(subtitle or "")
        subtitle_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle_paragraph.runs[0].font.size = Pt(11)
        for section_data in sections or []:
            document.add_heading(str(section_data.get("heading") or "Section"), level=1)
            if section_data.get("content"):
                document.add_paragraph(str(section_data["content"]))
            if section_data.get("callout"):
                paragraph = document.add_paragraph()
                run = paragraph.add_run(str(section_data["callout"]))
                run.bold = True
            table_data = section_data.get("table")
            if table_data:
                headers = list(table_data.get("headers") or [])
                rows = list(table_data.get("rows") or [])
                table = document.add_table(rows=1, cols=max(len(headers), 1))
                table.style = "Light Shading Accent 1"
                for cell, value in zip(table.rows[0].cells, headers):
                    cell.text = str(value)
                for row in rows:
                    cells = table.add_row().cells
                    for cell, value in zip(cells, row):
                        cell.text = str(value)
        document.save(output)
        return self._relative(output)

    def generate_pdf(self, filepath: str, title: str, subtitle: str, sections: list[dict[str, Any]]) -> str:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_CENTER
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import inch
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

        output = self._resolve_path(filepath, allowed_extensions=(".pdf",))
        styles = getSampleStyleSheet()
        styles.add(ParagraphStyle(name="ReportTitle", parent=styles["Title"], alignment=TA_CENTER, textColor=colors.HexColor("#163B52")))
        styles.add(ParagraphStyle(name="ReportSubtitle", parent=styles["Normal"], alignment=TA_CENTER, textColor=colors.HexColor("#57727D")))
        story = [Paragraph(title, styles["ReportTitle"]), Paragraph(subtitle or "", styles["ReportSubtitle"]), Spacer(1, 0.25 * inch)]
        for section_data in sections or []:
            story.append(Paragraph(str(section_data.get("heading") or "Section"), styles["Heading2"]))
            if section_data.get("content"):
                story.append(Paragraph(str(section_data["content"]), styles["BodyText"]))
            if section_data.get("callout"):
                callout = Table([[Paragraph(str(section_data["callout"]), styles["BodyText"])]], colWidths=[6.5 * inch])
                callout.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#E8F4F6")), ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#70B4B8")), ("LEFTPADDING", (0, 0), (-1, -1), 8), ("RIGHTPADDING", (0, 0), (-1, -1), 8)]))
                story.extend([Spacer(1, 0.1 * inch), callout])
            table_data = section_data.get("table")
            if table_data:
                rows = [list(table_data.get("headers") or [])] + [list(row) for row in table_data.get("rows") or []]
                if rows and rows[0]:
                    table = Table(rows, repeatRows=1)
                    table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#163B52")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#B7C9CC")), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
                    story.extend([Spacer(1, 0.12 * inch), table])
            story.append(Spacer(1, 0.18 * inch))

        def footer(canvas, document_obj):
            canvas.saveState()
            canvas.setFont("Helvetica", 8)
            canvas.setFillColor(colors.HexColor("#57727D"))
            canvas.drawRightString(7.7 * inch, 0.4 * inch, f"Page {document_obj.page}")
            canvas.restoreState()

        SimpleDocTemplate(str(output), pagesize=letter, rightMargin=0.7 * inch, leftMargin=0.7 * inch, topMargin=0.65 * inch, bottomMargin=0.7 * inch).build(story, onFirstPage=footer, onLaterPages=footer)
        return self._relative(output)

    def generate_note(self, filepath: str, title: str, tags: list[str], summary: str, sections: list[dict[str, Any]]) -> str:
        output = self._resolve_path(filepath, allowed_extensions=(".md", ".txt", ".markdown"))
        lines = ["---", f"title: {title}", f"tags: [{', '.join(tags)}]", "---", "", f"> {summary}", ""]
        for section_data in sections or []:
            lines.extend([f"## {section_data.get('heading', 'Section')}", ""])
            for task in section_data.get("tasks") or []:
                lines.append(f"- [ ] {task}")
            for bullet in section_data.get("bullets") or []:
                lines.append(f"- {bullet}")
            lines.append("")
        output.write_text("\n".join(lines), encoding="utf-8")
        return self._relative(output)

