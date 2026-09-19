"""Native XLSX parser using openpyxl, preserving spreadsheet structure, formulas, and merged ranges."""
from __future__ import annotations

import io
import logging
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, BinaryIO, Sequence
import openpyxl

from app.evidence.builder import EvidenceBuilder
from app.evidence.schema import EvidenceElement
from app.parsers.base import BaseParser

logger = logging.getLogger(__name__)


@dataclass
class XlsxCell:
    """Detailed representation of a spreadsheet cell distinguishing normal values, formulas, and empty states."""
    coordinate: str
    row: int
    column: int
    value: Any
    value_str: str
    is_formula: bool
    formula: str | None
    is_empty: bool
    data_type: str


@dataclass
class XlsxRow:
    """Represents a row of cells preserving column order."""
    row_index: int
    cells: list[XlsxCell] = field(default_factory=list)


@dataclass
class XlsxWorksheet:
    """Structured representation of a single worksheet."""
    name: str
    sheet_index: int
    max_row: int
    max_column: int
    merged_cells: list[str] = field(default_factory=list)
    rows: list[XlsxRow] = field(default_factory=list)
    formula_cells: list[XlsxCell] = field(default_factory=list)
    headers: list[str] = field(default_factory=list)
    data_rows: list[list[str]] = field(default_factory=list)

    @property
    def text(self) -> str:
        """Markdown formatted representation of the worksheet grid."""
        if not self.headers and not self.data_rows:
            return f"Sheet: {self.name} (empty)"
        header_line = " | ".join(self.headers)
        sep_line = " | ".join(["---"] * len(self.headers))
        body_lines = [" | ".join(r) for r in self.data_rows]
        return f"Sheet: {self.name}\n{header_line}\n{sep_line}\n" + "\n".join(body_lines)


@dataclass
class XlsxParseResult:
    """Structured result of Excel workbook parsing across all worksheets."""
    parser_name: str = "xlsx_parser"
    source_document: str = ""
    worksheets: list[XlsxWorksheet] = field(default_factory=list)
    metadata: dict[str, Any] = field(default_factory=dict)
    total_pages: int = 1
    parser_used: str = "openpyxl"

    _cached_elements: list[EvidenceElement] | None = field(default=None, repr=False)
    _doc_id: str = field(default="doc", repr=False)
    _hash: str = field(default="", repr=False)

    def to_evidence_elements(
        self,
        document_id: str | None = None,
        file_hash: str | None = None,
    ) -> list[EvidenceElement]:
        """Convert structured worksheet data and formulas into canonical EvidenceElements."""
        doc_id = document_id or self._doc_id
        f_hash = file_hash or self._hash
        elements: list[EvidenceElement] = []
        elem_idx = 1

        for ws in self.worksheets:
            # 1. Worksheet data table element
            if ws.headers or ws.data_rows:
                elem = EvidenceBuilder.build_element(
                    element_idx=elem_idx,
                    document_id=doc_id,
                    page=ws.sheet_index,
                    elem_type="table",
                    text=ws.text,
                    table_data={"headers": ws.headers, "rows": ws.data_rows},
                    bbox=[0.0, 0.0, 0.0, 0.0],
                    confidence=1.0,
                    extraction_model="openpyxl",
                    file_hash=f_hash,
                    sheet=ws.name,
                    metadata={
                        "sheet": ws.name,
                        "num_rows": len(ws.rows),
                        "num_cols": ws.max_column,
                        "merged_ranges": ws.merged_cells,
                    },
                )
                elements.append(elem)
                elem_idx += 1

            # 2. Distinct formula elements
            for fc in ws.formula_cells:
                elem = EvidenceBuilder.build_element(
                    element_idx=elem_idx,
                    document_id=doc_id,
                    page=ws.sheet_index,
                    elem_type="formula",
                    text=fc.formula,
                    formula_latex=fc.formula,
                    bbox=[0.0, 0.0, 0.0, 0.0],
                    confidence=1.0,
                    extraction_model="openpyxl",
                    file_hash=f_hash,
                    sheet=ws.name,
                    row=fc.row,
                    column=fc.column,
                    metadata={
                        "sheet": ws.name,
                        "cell_coordinate": fc.coordinate,
                        "is_formula": True,
                    },
                )
                elements.append(elem)
                elem_idx += 1

        return elements

    @property
    def elements(self) -> list[EvidenceElement]:
        """Convenience property for pipeline compatibility."""
        if self._cached_elements is None:
            self._cached_elements = self.to_evidence_elements()
        return self._cached_elements


class XlsxParser(BaseParser):
    """Native parser for Excel spreadsheets preserving cell formulas, merged ranges, and grid ordering."""

    @classmethod
    def _read_bytes(cls, input_data: bytes | BinaryIO | str | Path) -> bytes:
        if hasattr(input_data, "read"):
            return input_data.read()
        if isinstance(input_data, (str, Path)):
            with open(input_data, "rb") as f:
                return f.read()
        if isinstance(input_data, bytes):
            return input_data
        if isinstance(input_data, (bytearray, memoryview)):
            return bytes(input_data)
        raise TypeError(f"Expected bytes, Path, or file stream, got {type(input_data).__name__}")

    async def parse(
        self,
        file_bytes: bytes | BinaryIO | str | Path,
        filename: str = "",
        document_id: str = "doc",
        file_hash: str = "",
        **kwargs: Any,
    ) -> XlsxParseResult:
        """Parse an XLSX spreadsheet into structured worksheets, distinguishing values and formulas."""
        raw_bytes = self._read_bytes(file_bytes)
        source_doc = filename or "spreadsheet.xlsx"

        if not raw_bytes:
            return XlsxParseResult(
                parser_name="xlsx_parser",
                source_document=source_doc,
                worksheets=[],
                metadata={"total_sheets": 0, "is_empty": True},
                total_pages=1,
                parser_used="openpyxl",
                _doc_id=document_id,
                _hash=file_hash,
            )

        # data_only=False ensures formulas like =A1+B1 are preserved verbatim as strings
        wb = openpyxl.load_workbook(io.BytesIO(raw_bytes), data_only=False)
        total_sheets = len(wb.sheetnames)
        worksheet_objects: list[XlsxWorksheet] = []

        total_formulas = 0
        total_merged_ranges = 0

        for sheet_idx, sheet_name in enumerate(wb.sheetnames, start=1):
            ws = wb[sheet_name]
            merged_ranges = [str(r) for r in ws.merged_cells.ranges]
            total_merged_ranges += len(merged_ranges)

            rows: list[XlsxRow] = []
            formula_cells: list[XlsxCell] = []
            raw_grid_rows: list[list[str]] = []

            for row in ws.iter_rows(values_only=False):
                cell_records: list[XlsxCell] = []
                grid_row_vals: list[str] = []
                has_non_empty_cell = False

                for cell in row:
                    val = cell.value
                    data_type = str(cell.data_type)

                    # Empty cell check
                    is_empty = val is None or (isinstance(val, str) and not val.strip())

                    # Formula check
                    is_formula = False
                    formula_str = None
                    if isinstance(val, str) and val.startswith("="):
                        is_formula = True
                        formula_str = val
                    elif data_type == "f" and val is not None:
                        is_formula = True
                        formula_str = str(val) if str(val).startswith("=") else f"={val}"

                    if is_formula:
                        val_str = formula_str or ""
                        has_non_empty_cell = True
                    elif is_empty:
                        val_str = ""
                    else:
                        val_str = str(val).strip()
                        has_non_empty_cell = True

                    cell_record = XlsxCell(
                        coordinate=cell.coordinate,
                        row=cell.row,
                        column=cell.column,
                        value=val,
                        value_str=val_str,
                        is_formula=is_formula,
                        formula=formula_str,
                        is_empty=is_empty,
                        data_type=data_type,
                    )
                    cell_records.append(cell_record)
                    grid_row_vals.append(val_str)

                    if is_formula:
                        formula_cells.append(cell_record)
                        total_formulas += 1

                if has_non_empty_cell:
                    rows.append(XlsxRow(row_index=row[0].row if row else len(rows) + 1, cells=cell_records))
                    raw_grid_rows.append(grid_row_vals)

            headers = raw_grid_rows[0] if raw_grid_rows else []
            body_rows = raw_grid_rows[1:] if len(raw_grid_rows) > 1 else []

            ws_obj = XlsxWorksheet(
                name=sheet_name,
                sheet_index=sheet_idx,
                max_row=ws.max_row or 0,
                max_column=ws.max_column or 0,
                merged_cells=merged_ranges,
                rows=rows,
                formula_cells=formula_cells,
                headers=headers,
                data_rows=body_rows,
            )
            worksheet_objects.append(ws_obj)

        wb.close()

        meta = {
            "source_document": source_doc,
            "sheet_names": wb.sheetnames,
            "total_sheets": total_sheets,
            "total_formulas": total_formulas,
            "total_merged_ranges": total_merged_ranges,
            "is_empty": total_sheets == 0 or all(len(w.rows) == 0 for w in worksheet_objects),
        }

        return XlsxParseResult(
            parser_name="xlsx_parser",
            source_document=source_doc,
            worksheets=worksheet_objects,
            metadata=meta,
            total_pages=max(1, total_sheets),
            parser_used="openpyxl",
            _doc_id=document_id,
            _hash=file_hash,
        )


xlsx_parser = XlsxParser()
